"""Shared fixtures: tiny generated lecture decks + a scripted fake model (no API key needed)."""

from __future__ import annotations

from pathlib import Path

import pymupdf as fitz
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.models import Answer, Claim, Deck, Slide, Summary, VisualSpec
from app.store import Store

SORTING_SLIDES = [
    "Sorting Algorithms\nCS 2 - Lecture 7",
    "Merge Sort\nDivide and conquer\nRuns in O(n log n) time\nUses `merge()` to combine halves",
    "Quick Sort\nAverage case O(n log n)\nWorst case O(n^2) when the pivot is bad",
    "Benchmark\nSorting 1,000,000 integers took 0.42 seconds with merge sort",
]


@pytest.fixture
def sorting_deck() -> Deck:
    return Deck(
        id="d1",
        course="CS 2",
        title="Lecture 7 - Sorting",
        slides=[Slide(number=i + 1, text=t) for i, t in enumerate(SORTING_SLIDES)],
    )


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    path = tmp_path / "lecture7.pdf"
    doc = fitz.open()
    for text in SORTING_SLIDES:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(path)
    return path


@pytest.fixture
def pptx_file(tmp_path: Path) -> Path:
    path = tmp_path / "lecture8.pptx"
    prs = Presentation()
    for title, body, notes in [
        ("Hash Tables", "Average lookup is O(1)", "Collisions are handled with chaining"),
        ("Load Factor", "Resize when load factor exceeds 0.75", ""),
    ]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Comparison"
    table = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for (r, c), v in {(0, 0): "Structure", (0, 1): "Lookup", (1, 0): "Array", (1, 1): "O(n)"}.items():
        table.cell(r, c).text = v
    prs.save(path)
    return path


@pytest.fixture
def store() -> Store:
    return Store(":memory:")


GOOD_SUMMARY = Summary(
    abstract=[
        Claim(text="Merge sort is a divide and conquer algorithm that runs in O(n log n) time.", slides=[2]),
        Claim(text="Quick sort degrades to O(n^2) when the pivot is bad.", slides=[3]),
        Claim(text="Merge sort sorted 1,000,000 integers in 0.42 seconds.", slides=[4]),
    ],
    visuals=[
        VisualSpec(
            kind="table",
            title="Sorting complexity",
            data={
                "columns": ["Algorithm", "Time"],
                "rows": [["Merge sort", "O(n log n)"], ["Quick sort", "O(n log n) avg"]],
            },
            slides=[2, 3],
        )
    ],
)

HALLUCINATED_SUMMARY = Summary(
    abstract=[
        Claim(text="Merge sort runs in O(n log n) time.", slides=[2]),
        Claim(text="Sorting 1,000,000 integers took 0.9 seconds.", slides=[4]),  # wrong number
        Claim(text="Heap sort is covered on slide 9.", slides=[9]),  # slide doesn't exist
    ],
)


class FakeLLM:
    """Returns scripted outputs in order and records the feedback it received."""

    def __init__(self, summaries=(), answers=()):
        self.summaries, self.answers = list(summaries), list(answers)
        self.feedback_seen: list[list] = []

    def summarize(self, deck, feedback):
        self.feedback_seen.append(feedback)
        return self.summaries.pop(0).model_copy(deep=True)

    def answer(self, question, slides, feedback):
        self.feedback_seen.append(feedback)
        self.last_context = slides
        return self.answers.pop(0).model_copy(deep=True)


class FakeJudge:
    """Rejects any claim containing a banned phrase; accepts everything else."""

    def __init__(self, reject_if_contains: str = "__never__"):
        self.bad = reject_if_contains.lower()

    def supports(self, claim, evidence):
        if self.bad in claim.lower():
            return False, "claim not supported by evidence"
        return True, "ok"


@pytest.fixture
def fake_llm_factory():
    return FakeLLM


__all__ = ["GOOD_SUMMARY", "HALLUCINATED_SUMMARY", "FakeLLM", "FakeJudge", "Answer"]
