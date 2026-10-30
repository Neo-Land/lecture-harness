"""Turn PDF / PPTX lecture files into numbered slides."""

from __future__ import annotations

import hashlib
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf as fitz
from pptx import Presentation

from .models import Deck, Slide

SUPPORTED = {".pdf", ".pptx"}
IMAGE_DPI = 110

log = logging.getLogger("parsing")


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def _clean(text: str) -> str:
    lines = [ln.strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def _slide(index: int, text: str) -> Slide:
    cleaned = _clean(text)
    return Slide(number=index + 1, text=cleaned, needs_ocr=not cleaned)


def parse_pdf(path: Path) -> list[Slide]:
    with fitz.open(path) as doc:
        return [_slide(i, page.get_text()) for i, page in enumerate(doc)]


def parse_pptx(path: Path) -> list[Slide]:
    slides = []
    for i, s in enumerate(Presentation(str(path)).slides):
        parts = []
        for shape in s.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if s.has_notes_slide:  # speaker notes often hold the real explanation
            parts.append(s.notes_slide.notes_text_frame.text)
        slides.append(_slide(i, "\n".join(parts)))
    return slides


def parse_file(path: Path, course: str, deck_id: str | None = None) -> Deck:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {ext}")
    if not path.is_file():
        raise FileNotFoundError(f"No such file: {path}")
    try:
        slides = parse_pdf(path) if ext == ".pdf" else parse_pptx(path)
    except Exception as err:  # the PDF and PPTX libraries raise many unrelated types on bad files
        raise ValueError(f"Could not read {path.name}: not a valid {ext} file") from err
    return Deck(id=deck_id or file_hash(path), course=course, title=path.stem, slides=slides)


def _render_pdf(path: Path, out_dir: Path) -> set[int]:
    out_dir.mkdir(parents=True, exist_ok=True)
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            page.get_pixmap(dpi=IMAGE_DPI).save(out_dir / f"{i + 1}.png")
        return set(range(1, len(doc) + 1))


def render_images(path: Path, out_dir: Path) -> set[int]:
    """Render slides to <out_dir>/<n>.png and return the numbers rendered. Best effort: a failure
    means text-only slides, never a failed ingest. PPTX needs LibreOffice to convert first."""
    try:
        if path.suffix.lower() == ".pdf":
            return _render_pdf(path, out_dir)
        soffice = shutil.which("soffice")
        if not soffice:
            log.info("no_libreoffice", extra={"file": path.name})
            return set()
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
                check=True,
                capture_output=True,
                timeout=180,
            )
            return _render_pdf(Path(tmp) / f"{path.stem}.pdf", out_dir)
    except Exception as err:
        log.warning("render_failed", extra={"file": path.name, "error": type(err).__name__})
        return set()
